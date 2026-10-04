#!/usr/bin/env python3
"""Aplica resumen oficial MP cuenta sept 2026 sobre CIUDAD Excel."""
from __future__ import annotations

import json
import math
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path("/workspace")
XLSX = ROOT / "CIUDAD/Gastos_CIUDAD_1132_2026.xlsx"
JSON = ROOT / "CIUDAD/_inbox/septiembre_2026/mp_exports/mp_cuenta_oficial_sept.json"
SRC = "resumen oficial cuenta MP 1–30 sep 2026"


def ceil_ars(v: float) -> int:
    return int(math.ceil(abs(float(v))))


def parse_fecha(s: str) -> datetime:
    return datetime.strptime(s, "%Y-%m-%d")


def total_formula(row: int) -> str:
    return (
        f'=IF(OR(D{row}<>"",E{row}<>""),'
        f'IF(D{row}="",0,D{row})+IF(E{row}="",0,E{row}),"")'
    )


def clear_row(ws, row: int, cols: int = 8) -> None:
    for c in range(1, cols + 1):
        ws.cell(row, c).value = None


def next_empty_row(ws, start: int, key_col: int = 1) -> int:
    r = start
    while ws.cell(r, key_col).value not in (None, ""):
        r += 1
    return r


def expand_sumif(ws, rows: range, new_end: int = 500) -> None:
    for r in rows:
        for c in range(1, 8):
            val = ws.cell(r, c).value
            if isinstance(val, str) and "SUMIF" in val:
                ws.cell(r, c).value = re.sub(
                    r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$\d+",
                    lambda m: f"${m.group(1)}${m.group(2)}:${m.group(3)}${new_end}",
                    val,
                )


def main() -> None:
    txns = json.loads(JSON.read_text())
    wb = load_workbook(XLSX)

    gps = wb["Gasto Personal Sol"]
    alm = wb["Almacen Detalle"]
    mpng = wb["MP No Gasto"]
    inv = wb["Inversión"]
    tr = wb["Transferencias recibidas"]
    dr = wb["Dinero Retirado"]
    pt = wb["Pago de Tarjetas"]
    usd = wb["Sol Dollar Mdp"]

    # --- 1) Anular cargos con devolución neta 0 ---
    # GPS: Didi 7600 (2/9), un Ebanx 6607 (2/9), Ebanx 6053 (18/9)
    anulados_gps = []
    ebanx_6607_cleared = False
    for r in range(26, gps.max_row + 1):
        f = gps.cell(r, 1).value
        if not isinstance(f, datetime) or f.year != 2026 or f.month != 9:
            continue
        concepto = str(gps.cell(r, 3).value or "")
        monto = gps.cell(r, 4).value
        medio = str(gps.cell(r, 5).value or "")
        if medio != "Mercado Pago":
            continue
        if f.day == 2 and concepto == "TRANSPORTE · Didi" and monto == 7600:
            clear_row(gps, r, 6)
            anulados_gps.append((r, "Didi 7600"))
        elif (
            f.day == 2
            and concepto == "TRANSPORTE · Uber"
            and monto == 6607
            and not ebanx_6607_cleared
        ):
            clear_row(gps, r, 6)
            ebanx_6607_cleared = True
            anulados_gps.append((r, "Uber/Ebanx 6607 (devuelto)"))
        elif f.day == 18 and concepto == "TRANSPORTE · Uber" and monto == 6053:
            clear_row(gps, r, 6)
            anulados_gps.append((r, "Uber/Ebanx 6053"))

    # Almacén: PedidosYa 38099 (12/9) anulado
    anulados_alm = []
    for r in range(4, alm.max_row + 1):
        f = alm.cell(r, 2).value
        if not isinstance(f, datetime) or f.year != 2026 or f.month != 9:
            continue
        if (
            f.day == 12
            and alm.cell(r, 3).value == "PEDIDOS YA · COMIDA"
            and alm.cell(r, 4).value == 38099
        ):
            clear_row(alm, r, 7)
            # dejar mes vacío también
            anulados_alm.append(r)

    # --- 2) Nuevos gastos 28–30 ---
    new_alm = [
        (datetime(2026, 9, 28), "COMIDA · pollo", 10900, "COMIDA · pollo · Súper Pollo · QR · " + SRC),
        (datetime(2026, 9, 30), "Combustible — Shell", 1039, "Combustible — Shell · QR · " + SRC),
        (datetime(2026, 9, 30), "Combustible — Shell", 288, "Combustible — Shell · QR · ceil 287.50 · " + SRC),
        (datetime(2026, 9, 30), "Combustible — Shell", 1126, "Combustible — Shell · QR · " + SRC),
        (datetime(2026, 9, 30), "Combustible — Shell", 62548, "Combustible — Shell · QR · ceil 62547.50 · " + SRC),
    ]
    # usar filas vacías de septiembre (381+)
    slot = 381
    for fecha, detalle, monto, comentario in new_alm:
        while slot < 396 and alm.cell(slot, 3).value not in (None, ""):
            slot += 1
        if slot >= 396:
            raise RuntimeError("Sin slots vacíos en Almacén septiembre")
        alm.cell(slot, 1).value = "Septiembre"
        alm.cell(slot, 2).value = fecha
        alm.cell(slot, 3).value = detalle
        alm.cell(slot, 4).value = monto
        alm.cell(slot, 5).value = None
        alm.cell(slot, 6).value = total_formula(slot)
        alm.cell(slot, 7).value = comentario
        slot += 1

    new_gps = [
        (datetime(2026, 9, 28), "TRANSFERENCIAS", 69900, "Mercado Pago",
         "TRANSFERENCIAS · Camilo Carlos Ihan · REVISAR ítem · " + SRC),
        (datetime(2026, 9, 29), "TRANSPORTE · Uber", 7230, "Mercado Pago",
         "TRANSPORTE · Uber · Ebanx (=Uber) · " + SRC),
        (datetime(2026, 9, 29), "TRANSPORTE · Didi", 5700, "Mercado Pago",
         "TRANSPORTE · Didi · " + SRC),
        (datetime(2026, 9, 29), "TRANSPORTE · Didi", 4000, "Mercado Pago",
         "TRANSPORTE · Didi · " + SRC),
        (datetime(2026, 9, 30), "TRANSFERENCIAS", 20000, "Mercado Pago",
         "TRANSFERENCIAS · Kelly · REVISAR ítem · " + SRC),
        (datetime(2026, 9, 30), "TRANSFERENCIAS", 86000, "Mercado Pago",
         "TRANSFERENCIAS · Kelly · REVISAR ítem · " + SRC),
    ]
    gps_row = next_empty_row(gps, 938)
    for fecha, concepto, monto, medio, comentario in new_gps:
        gps.cell(gps_row, 1).value = fecha
        gps.cell(gps_row, 2).value = "Septiembre"
        gps.cell(gps_row, 3).value = concepto
        gps.cell(gps_row, 4).value = monto
        gps.cell(gps_row, 5).value = medio
        gps.cell(gps_row, 6).value = comentario
        gps_row += 1

    # --- 3) MP No Gasto + espejos ---
    # Evitar duplicar si ya hay sept
    existing_mpng_sept = set()
    for r in range(24, mpng.max_row + 1):
        f = mpng.cell(r, 1).value
        if isinstance(f, datetime) and f.year == 2026 and f.month == 9:
            key = (
                f.date().isoformat(),
                str(mpng.cell(r, 3).value),
                str(mpng.cell(r, 4).value),
                mpng.cell(r, 5).value,
            )
            existing_mpng_sept.add(key)

    mpng_row = next_empty_row(mpng, 54)
    dr_row = next_empty_row(dr, 33)
    pt_row = next_empty_row(pt, 30)
    inv_row = next_empty_row(inv, 26)
    tr_row = next_empty_row(tr, 44)

    expand_sumif(tr, range(6, 18), 500)

    stats = {
        "mpng": 0,
        "dr": 0,
        "pt": 0,
        "inv": 0,
        "tr": 0,
        "skip_dup": 0,
    }

    for t in txns:
        desc = t["desc"]
        dlow = desc.lower()
        fecha = parse_fecha(t["fecha"])
        valor = float(t["valor"])
        abs_m = ceil_ars(valor)
        op = t.get("op") or ""

        # Dinero retirado (ARS pots) — no "Retiro Bonos"
        if dlow.startswith("dinero retirado") and "bonos" not in dlow:
            concepto = desc.replace("Dinero retirado", "").strip() or "Gastos"
            key = (fecha.date().isoformat(), "Dinero retirado", concepto, abs_m)
            if key in existing_mpng_sept:
                stats["skip_dup"] += 1
                continue
            mpng.cell(mpng_row, 1).value = fecha
            mpng.cell(mpng_row, 2).value = "Septiembre"
            mpng.cell(mpng_row, 3).value = "Dinero retirado"
            mpng.cell(mpng_row, 4).value = concepto
            mpng.cell(mpng_row, 5).value = abs_m
            mpng.cell(mpng_row, 6).value = "ARS"
            mpng.cell(mpng_row, 7).value = "Mercado Pago"
            mpng.cell(mpng_row, 8).value = f"Dinero retirado · + · op {op} · {SRC}"
            mpng_row += 1
            stats["mpng"] += 1

            dr.cell(dr_row, 1).value = fecha
            dr.cell(dr_row, 2).value = "Septiembre"
            dr.cell(dr_row, 3).value = concepto
            dr.cell(dr_row, 4).value = abs_m
            dr.cell(dr_row, 5).value = "Mercado Pago · ARS"
            dr.cell(dr_row, 6).value = f"Dinero retirado · + · op {op} · {SRC}"
            dr_row += 1
            stats["dr"] += 1
            continue

        if dlow.startswith("reserva por gastos"):
            concepto = desc.replace("Reserva por gastos", "").strip()
            key = (fecha.date().isoformat(), "Reserva automática", concepto, abs_m)
            if key in existing_mpng_sept:
                stats["skip_dup"] += 1
                continue
            # reservas diarias repetidas: permitir varias del mismo monto/día
            mpng.cell(mpng_row, 1).value = fecha
            mpng.cell(mpng_row, 2).value = "Septiembre"
            mpng.cell(mpng_row, 3).value = "Reserva automática"
            mpng.cell(mpng_row, 4).value = concepto
            mpng.cell(mpng_row, 5).value = abs_m
            mpng.cell(mpng_row, 6).value = "ARS"
            mpng.cell(mpng_row, 7).value = "Mercado Pago"
            mpng.cell(mpng_row, 8).value = f"Reserva automática · − · op {op} · NUNCA almacén · {SRC}"
            mpng_row += 1
            stats["mpng"] += 1
            continue

        if dlow.startswith("dinero reservado"):
            concepto = desc.replace("Dinero reservado", "").strip()
            key = (fecha.date().isoformat(), "Dinero reservado", concepto, abs_m)
            if key in existing_mpng_sept:
                stats["skip_dup"] += 1
                continue
            mpng.cell(mpng_row, 1).value = fecha
            mpng.cell(mpng_row, 2).value = "Septiembre"
            mpng.cell(mpng_row, 3).value = "Dinero reservado"
            mpng.cell(mpng_row, 4).value = concepto
            mpng.cell(mpng_row, 5).value = abs_m
            mpng.cell(mpng_row, 6).value = "ARS"
            mpng.cell(mpng_row, 7).value = "Mercado Pago"
            mpng.cell(mpng_row, 8).value = f"Dinero reservado · − · op {op} · NUNCA almacén · {SRC}"
            mpng_row += 1
            stats["mpng"] += 1
            continue

        if "pago de resumen" in dlow:
            mpng.cell(mpng_row, 1).value = fecha
            mpng.cell(mpng_row, 2).value = "Septiembre"
            mpng.cell(mpng_row, 3).value = "Pago de resumen TC"
            mpng.cell(mpng_row, 4).value = "Tarjeta de crédito Mercado Pago"
            mpng.cell(mpng_row, 5).value = abs_m
            mpng.cell(mpng_row, 6).value = "ARS"
            mpng.cell(mpng_row, 7).value = "Mercado Pago"
            mpng.cell(mpng_row, 8).value = f"Pago de resumen · − · op {op} · NUNCA almacén · {SRC}"
            mpng_row += 1
            stats["mpng"] += 1

            pt.cell(pt_row, 1).value = fecha
            pt.cell(pt_row, 2).value = "Septiembre"
            pt.cell(pt_row, 3).value = "Tarjeta de crédito"
            pt.cell(pt_row, 4).value = abs_m
            pt.cell(pt_row, 5).value = "Mercado Pago"
            pt.cell(pt_row, 6).value = f"Pago de resumen · ${abs_m:,} · op {op} · {SRC}".replace(",", ".")
            pt_row += 1
            stats["pt"] += 1
            continue

        # Inversión ARS (créditos de venta / retiro bonos) — no líneas US$ del JSON ARS
        if valor > 0 and (
            "venta de dólar" in dlow
            or "venta de dolar" in dlow
            or "retiro bonos" in dlow
        ):
            if "retiro bonos" in dlow:
                tipo = "Bonos/PF"
                instrumento = "Bonos, plazos fijos y más"
                concepto_mp = "Bonos, plazos fijos y más"
            else:
                tipo = "USD"
                instrumento = "Venta de dólares"
                concepto_mp = "Venta de dólares"

            mpng.cell(mpng_row, 1).value = fecha
            mpng.cell(mpng_row, 2).value = "Septiembre"
            mpng.cell(mpng_row, 3).value = "Inversión"
            mpng.cell(mpng_row, 4).value = concepto_mp
            mpng.cell(mpng_row, 5).value = abs_m
            mpng.cell(mpng_row, 6).value = "ARS"
            mpng.cell(mpng_row, 7).value = "Mercado Pago"
            mpng.cell(mpng_row, 8).value = f"Inversión · Rescate · + · op {op} · ceil {valor} · {SRC}"
            mpng_row += 1
            stats["mpng"] += 1

            inv.cell(inv_row, 1).value = fecha
            inv.cell(inv_row, 2).value = "Septiembre"
            inv.cell(inv_row, 3).value = tipo
            inv.cell(inv_row, 4).value = "Rescate"
            inv.cell(inv_row, 5).value = instrumento
            inv.cell(inv_row, 6).value = abs_m
            inv.cell(inv_row, 7).value = "Mercado Pago"
            inv.cell(inv_row, 8).value = None
            inv.cell(inv_row, 9).value = None
            inv.cell(inv_row, 10).value = f"Rescate · op {op} · {SRC}"
            inv_row += 1
            stats["inv"] += 1
            continue

        if dlow.startswith("transferencia recibida"):
            origen = desc.replace("Transferencia recibida", "").strip()
            tr.cell(tr_row, 1).value = fecha
            tr.cell(tr_row, 2).value = "Septiembre"
            tr.cell(tr_row, 3).value = origen
            tr.cell(tr_row, 4).value = abs_m
            tr.cell(tr_row, 5).value = "Mercado Pago — Transferencia recibida"
            tr.cell(tr_row, 6).value = f"recibida · op {op} · {SRC}"
            tr_row += 1
            stats["tr"] += 1
            continue

    # --- 4) Sol Dollar Mdp desde tenencias oficiales ---
    # Header oficial: saldo final US$ 1,68. Detalle PDF cierra en 2002,21 (inconsistencia del extracto).
    # Prevalece el saldo final del resumen + confirmación Sol (~0). Ajuste REVISAR por el gap.
    usd["A4"] = "Saldo USD a la fecha"
    usd["B4"] = datetime(2026, 9, 30)
    usd["C4"] = 1.68
    usd["D4"] = "USD"
    usd["E4"] = (
        "Resumen oficial tenencias USD MP · saldo final 1–30 sep = US$ 1,68 "
        "(detalle de movimientos del PDF cierra en 2002,21 → ajuste REVISAR)"
    )

    usd_moves = [
        (datetime(2026, 9, 1), "Saldo inicial", "Tenencias MP", "Stock USD inicio septiembre", 3159.41, None, None, 3159.41,
         "Saldo inicial oficial tenencias USD"),
        (datetime(2026, 9, 1), "Rendimiento", "Mercado Pago", "Rendimientos USD", 0.06, None, None, 3159.47,
         "Tenencias USD · " + SRC),
        (datetime(2026, 9, 1), "Venta", "Mercado Pago", "Venta de dólar MEP", -1135.23, 1511.25, 1715612, 2024.24,
         "Venta MEP · US$ -1.135,23 · ARS ceil 1.715.611,40 · " + SRC),
        (datetime(2026, 9, 7), "Entrada", "Mercado Pago", "Dinero retirado Gastos (USD)", 2001.20, None, None, 4025.44,
         "Dinero retirado Gastos US$ 2.001,20 · " + SRC),
        (datetime(2026, 9, 8), "Rendimiento", "Mercado Pago", "Rendimientos USD", 0.10, None, None, 4025.54,
         "Tenencias USD · " + SRC),
        (datetime(2026, 9, 9), "Rendimiento", "Mercado Pago", "Rendimientos USD", 0.10, None, None, 4025.64,
         "Tenencias USD · " + SRC),
        (datetime(2026, 9, 10), "Rendimiento", "Mercado Pago", "Rendimientos USD", 0.10, None, None, 4025.74,
         "Tenencias USD · " + SRC),
        (datetime(2026, 9, 10), "Venta", "Mercado Pago", "Venta de dólar MEP", -1984.17, 1506.78, 2989714, 2041.57,
         "Venta MEP · US$ -1.984,17 · ARS ceil 2.989.713,19 · " + SRC),
        (datetime(2026, 9, 15), "Rendimiento", "Mercado Pago", "Rendimientos USD", 0.01, None, None, 2041.58,
         "Tenencias USD · " + SRC),
        (datetime(2026, 9, 15), "Venta", "Mercado Pago", "Venta de dólar MEP", -39.37, 1507.08, 59334, 2002.21,
         "Venta MEP · US$ -39,37 · ARS ceil 59.333,83 · " + SRC),
        (datetime(2026, 9, 30), "Ajuste", "Tenencias MP", "Ajuste a saldo final oficial", -2000.53, None, None, 1.68,
         "REVISAR · detalle PDF termina en US$ 2002,21 pero saldo final del resumen = US$ 1,68 · " + SRC),
    ]
    # Reemplazar fila 8 (saldo inicial viejo) e insertar el resto antes del bloque consumos MC
    clear_row(usd, 8, 9)
    need = len(usd_moves)
    if need > 1:
        usd.insert_rows(9, amount=need - 1)

    for i, row in enumerate(usd_moves):
        r = 8 + i
        for c, v in enumerate(row, start=1):
            usd.cell(r, c).value = v

    header_row = None
    for r in range(1, 40):
        if usd.cell(r, 1).value == "Fecha" and usd.cell(r, 2).value == "Tipo":
            header_row = r
            break
    assert header_row == 7, header_row

    wb.save(XLSX)

    # Totales almacén sept post-cambio
    alm_sum = 0
    for r in range(4, alm.max_row + 1):
        if alm.cell(r, 1).value == "Septiembre" and isinstance(alm.cell(r, 4).value, (int, float)):
            alm_sum += alm.cell(r, 4).value

    gps_sum = 0
    gps_n = 0
    for r in range(26, gps.max_row + 1):
        if gps.cell(r, 2).value == "Septiembre" and isinstance(gps.cell(r, 4).value, (int, float)):
            gps_sum += gps.cell(r, 4).value
            gps_n += 1

    print("anulados_gps", anulados_gps)
    print("anulados_alm", anulados_alm)
    print("stats", stats)
    print("almacen_sept", alm_sum)
    print("gps_sept_rows", gps_n, "sum", gps_sum)
    print("saved", XLSX)


if __name__ == "__main__":
    main()
